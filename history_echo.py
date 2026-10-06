"""History Echo News: daily headlines + "has this happened before, and what came next?"

Pipeline: RSS headlines -> LLM picks key stories and finds historical precedents
-> precedents are checked against Wikipedia -> Markdown report, Excel board,
GitHub Issue (notification) and optional email.
"""

from __future__ import annotations

import argparse
import html
import json
import os
import re
import smtplib
import ssl
import sys
import time
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from pathlib import Path
from urllib.parse import quote
from zoneinfo import ZoneInfo

import feedparser
import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from feeds import CATEGORIES

ROOT = Path(__file__).resolve().parent
REPORTS = ROOT / "reports"
DATA = ROOT / "data"
BOARD = DATA / "Newspaper_Board.xlsx"
TZ = ZoneInfo("America/New_York")
UA = {"User-Agent": "HistoryEchoNews/1.0 (https://github.com/movlan-aliyev)"}

MAX_HEADLINES_PER_CATEGORY = 25
STORIES_PER_CATEGORY = int(os.environ.get("STORIES_PER_CATEGORY", "2"))
LOOKBACK_HOURS = 36


def fetch_headlines() -> dict[str, list[dict]]:
    cutoff = datetime.now(timezone.utc) - timedelta(hours=LOOKBACK_HOURS)
    out: dict[str, list[dict]] = {}
    for category, urls in CATEGORIES.items():
        seen: set[str] = set()
        items: list[dict] = []
        for url in urls:
            try:
                resp = requests.get(url, headers=UA, timeout=20)
                feed = feedparser.parse(resp.content)
            except Exception as exc:
                print(f"  ! feed failed {url}: {exc}")
                continue
            source = feed.feed.get("title", url)
            for e in feed.entries:
                title = (e.get("title") or "").strip()
                key = re.sub(r"\W+", "", title.lower())[:80]
                if not title or key in seen:
                    continue
                published = e.get("published_parsed") or e.get("updated_parsed")
                if published and datetime(*published[:6], tzinfo=timezone.utc) < cutoff:
                    continue
                seen.add(key)
                summary = re.sub(r"<[^>]+>", "", e.get("summary", "")).strip()
                items.append({
                    "title": title,
                    "summary": summary[:300],
                    "source": source,
                    "link": e.get("link", ""),
                })
        out[category] = items[:MAX_HEADLINES_PER_CATEGORY]
        print(f"  {category}: {len(out[category])} headlines")
    return out


def parse_json(text: str) -> dict:
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text.strip())
    start, end = text.find("{"), text.rfind("}")
    if start == -1 or end == -1:
        raise ValueError(f"no JSON object in model reply: {text[:200]!r}")
    return json.loads(text[start:end + 1])


class LLM:
    """GitHub Models (free via GITHUB_TOKEN in Actions) with OpenAI as optional fallback."""

    def __init__(self) -> None:
        self.openai_key = os.environ.get("OPENAI_API_KEY", "").strip()
        self.gh_token = os.environ.get("GITHUB_TOKEN", "").strip()
        self.models = [m for m in os.environ.get(
            "LLM_MODELS", "openai/gpt-4.1,openai/gpt-4.1-mini,openai/gpt-4o-mini").split(",") if m]

    def chat(self, system: str, user: str) -> dict:
        attempts: list[tuple[str, dict, str]] = []
        if self.gh_token:
            for m in self.models:
                attempts.append(("https://models.github.ai/inference/chat/completions",
                                 {"Authorization": f"Bearer {self.gh_token}"}, m))
        if self.openai_key:
            attempts.append(("https://api.openai.com/v1/chat/completions",
                             {"Authorization": f"Bearer {self.openai_key}"},
                             os.environ.get("OPENAI_MODEL", "gpt-4.1-mini")))
        if not attempts:
            raise RuntimeError("No GITHUB_TOKEN or OPENAI_API_KEY available for the LLM.")

        last_err = ""
        for url, headers, model in attempts:
            for retry in range(2):
                try:
                    r = requests.post(url, headers={**headers, "Content-Type": "application/json"},
                                      json={
                                          "model": model,
                                          "temperature": 0.3,
                                          "response_format": {"type": "json_object"},
                                          "messages": [{"role": "system", "content": system},
                                                       {"role": "user", "content": user}],
                                      }, timeout=180)
                    if r.status_code == 429:
                        last_err = f"{model}: rate limited"
                        time.sleep(20 * (retry + 1))
                        continue
                    if r.status_code >= 400 or not r.text.strip():
                        last_err = f"{model}: HTTP {r.status_code} {r.text[:300]!r}"
                        break
                    content = r.json()["choices"][0]["message"]["content"] or ""
                    print(f"    (model: {model})")
                    return parse_json(content)
                except Exception as exc:
                    last_err = f"{model}: {type(exc).__name__}: {exc}"
                    break
            print(f"    ! {last_err}; trying next model")
        raise RuntimeError(last_err)


SYSTEM_PROMPT = """You are "History Echo", an analyst who connects today's news to history.
For each story you choose, find the same or similar events from the past (days,
months, years or centuries ago), say what happened AFTER them, and use that track
record to estimate what is likely to happen next. Be factual: never invent events.

Analyst guidance from the user (follow it closely):
---
{guidance}
---

Respond ONLY with JSON in this exact shape:
{{"stories": [{{
  "headline": "short headline of today's story",
  "source_link": "link copied from the input",
  "what_happened": "1-2 sentences on today's news",
  "event_type": "e.g. ceasefire, rate cut, manager sacked, shutdown",
  "precedents": [{{
     "when": "year or date",
     "event": "name of the past event",
     "wikipedia_title": "exact English Wikipedia article title",
     "what_happened": "1 sentence",
     "what_followed": "what happened afterwards and how soon",
     "similarity": "why it is comparable (same actors / same type)"
  }}],
  "pattern": "the lesson history teaches, with a base rate if possible",
  "outlook": "what is likely to happen next",
  "probability": "e.g. 65%",
  "timeframe": "e.g. next 6 months",
  "confidence": "Low | Medium | High",
  "different_this_time": "factors that could break the pattern",
  "watch_for": "early signals to watch"
}}]}}"""


def analyze(llm: LLM, category: str, headlines: list[dict], guidance: str) -> list[dict]:
    if not headlines:
        return []
    listing = "\n".join(
        f"{i+1}. {h['title']} | {h['summary'][:200]} | {h['source']} | {h['link']}"
        for i, h in enumerate(headlines))
    user = (f"Category: {category}\nToday: {datetime.now(TZ):%A %d %B %Y}\n\n"
            f"Today's headlines:\n{listing}\n\n"
            f"Choose the {STORIES_PER_CATEGORY} most significant stories where history gives real "
            f"insight. Give 2-3 precedents each, strongest first.")
    data = llm.chat(SYSTEM_PROMPT.format(guidance=guidance), user)
    stories = data.get("stories", [])[:STORIES_PER_CATEGORY]
    for s in stories:
        s["category"] = category
    return stories


def verify_on_wikipedia(title: str) -> dict | None:
    if not title:
        return None
    try:
        r = requests.get("https://en.wikipedia.org/w/api.php", headers=UA, timeout=15, params={
            "action": "query", "list": "search", "srsearch": title, "srlimit": 1, "format": "json"})
        hits = r.json().get("query", {}).get("search", [])
        if not hits:
            return None
        found = hits[0]["title"]
        s = requests.get(
            f"https://en.wikipedia.org/api/rest_v1/page/summary/{quote(found.replace(' ', '_'))}",
            headers=UA, timeout=15).json()
        return {"title": found,
                "url": s.get("content_urls", {}).get("desktop", {}).get("page", ""),
                "extract": (s.get("extract") or "")[:280]}
    except Exception:
        return None


def verify(stories: list[dict]) -> None:
    for s in stories:
        for p in s.get("precedents", []):
            p["wiki"] = verify_on_wikipedia(p.get("wikipedia_title") or p.get("event", ""))


def render_markdown(day: datetime, stories: list[dict], headline_counts: dict[str, int]) -> str:
    lines = [f"# History Echo \u2014 {day:%A, %B %d, %Y}", "",
             "_Today's news, and what happened the last time something like it happened._", ""]
    for category in CATEGORIES:
        cat = [s for s in stories if s["category"] == category]
        if not cat:
            continue
        lines += [f"## {category}", ""]
        for s in cat:
            link = s.get("source_link", "")
            lines += [f"### {s.get('headline', '')}" + (f" ([source]({link}))" if link else ""), "",
                      f"**Today:** {s.get('what_happened', '')}", "",
                      "**Has it happened before?**", ""]
            for p in s.get("precedents", []):
                w = p.get("wiki")
                ref = f" \u2014 [Wikipedia: {w['title']}]({w['url']})" if w and w.get("url") else " \u2014 _(not verified)_"
                lines += [f"- **{p.get('when', '')} \u2014 {p.get('event', '')}**{ref}",
                          f"  - What happened: {p.get('what_happened', '')}",
                          f"  - What followed: {p.get('what_followed', '')}",
                          f"  - Why comparable: {p.get('similarity', '')}"]
            lines += ["",
                      f"**Pattern:** {s.get('pattern', '')}", "",
                      f"**Outlook:** {s.get('outlook', '')} \u2014 **{s.get('probability', '?')}** within "
                      f"{s.get('timeframe', '?')} (confidence: {s.get('confidence', '?')})", "",
                      f"**Why this time could be different:** {s.get('different_this_time', '')}", "",
                      f"**Watch for:** {s.get('watch_for', '')}", "", "---", ""]
    total = sum(headline_counts.values())
    lines += [f"_Scanned {total} headlines from free RSS feeds. Outlooks are pattern-based estimates, "
              f"not certainties. Edit `guidance.md` to teach the agent._"]
    return "\n".join(lines)


def render_html(md: str) -> str:
    out = []
    for line in md.splitlines():
        esc = html.escape(line)
        esc = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r'<a href="\2">\1</a>', esc)
        esc = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", esc)
        esc = re.sub(r"_(.+?)_", r"<i>\1</i>", esc)
        if line.startswith("### "):
            out.append(f"<h3 style='margin-bottom:4px'>{esc[4:]}</h3>")
        elif line.startswith("## "):
            out.append(f"<h2 style='color:#1f4e79;border-bottom:2px solid #1f4e79'>{esc[3:]}</h2>")
        elif line.startswith("# "):
            out.append(f"<h1>{esc[2:]}</h1>")
        elif line.startswith("  - "):
            out.append(f"<div style='margin-left:28px'>\u2022 {esc[4:]}</div>")
        elif line.startswith("- "):
            out.append(f"<div style='margin-left:12px;margin-top:6px'>\u25b8 {esc[2:]}</div>")
        elif line == "---":
            out.append("<hr>")
        elif line:
            out.append(f"<p style='margin:4px 0'>{esc}</p>")
    return ("<html><body style='font-family:Segoe UI,Arial,sans-serif;max-width:820px;"
            "line-height:1.45'>" + "\n".join(out) + "</body></html>")


BOARD_COLUMNS = ["Date", "Category", "Headline", "Today", "Event Type", "Historical Precedents",
                 "What Followed Before", "Pattern", "Outlook", "Probability", "Timeframe",
                 "Confidence", "Why Different This Time", "Watch For", "Source", "Verified Links"]
WIDTHS = [12, 22, 40, 50, 18, 50, 60, 50, 45, 12, 16, 12, 45, 40, 30, 50]


def update_board(path: Path, day: datetime, stories: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        wb = load_workbook(path)
        ws = wb["Daily Echoes"] if "Daily Echoes" in wb.sheetnames else wb.active
    else:
        wb = Workbook()
        ws = wb.active
    if ws.title != "Daily Echoes" or ws.max_row < 1 or ws.cell(1, 1).value != "Date":
        ws.title = "Daily Echoes"
        ws.delete_rows(1, ws.max_row)
        ws.append(BOARD_COLUMNS)
        for i, w in enumerate(WIDTHS, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
            c = ws.cell(1, i)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E79")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"

    today = day.strftime("%Y-%m-%d")
    for r in range(ws.max_row, 1, -1):
        if str(ws.cell(r, 1).value) == today:
            ws.delete_rows(r)

    for s in stories:
        precs = s.get("precedents", [])
        ws.append([
            today, s["category"], s.get("headline", ""), s.get("what_happened", ""),
            s.get("event_type", ""),
            "\n".join(f"{p.get('when', '')}: {p.get('event', '')}" for p in precs),
            "\n".join(f"{p.get('when', '')}: {p.get('what_followed', '')}" for p in precs),
            s.get("pattern", ""), s.get("outlook", ""), s.get("probability", ""),
            s.get("timeframe", ""), s.get("confidence", ""), s.get("different_this_time", ""),
            s.get("watch_for", ""), s.get("source_link", ""),
            "\n".join(p["wiki"]["url"] for p in precs if p.get("wiki") and p["wiki"].get("url")),
        ])
        for c in ws[ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


def send_email(subject: str, html_body: str, text_body: str) -> None:
    host = os.environ.get("SMTP_HOST", "").strip()
    user = os.environ.get("SMTP_USER", "").strip()
    pwd = os.environ.get("SMTP_PASSWORD", "").strip()
    to = os.environ.get("MAIL_TO", "").strip()
    if not all([host, user, pwd, to]):
        print("  Email skipped (SMTP secrets not set).")
        return
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = os.environ.get("MAIL_FROM", user).strip() or user
    msg["To"] = to
    msg.set_content(text_body)
    msg.add_alternative(html_body, subtype="html")
    port = int(os.environ.get("SMTP_PORT", "587"))
    ctx = ssl.create_default_context()
    if port == 465:
        with smtplib.SMTP_SSL(host, port, context=ctx) as s:
            s.login(user, pwd)
            s.send_message(msg)
    else:
        with smtplib.SMTP(host, port) as s:
            s.starttls(context=ctx)
            s.login(user, pwd)
            s.send_message(msg)
    print(f"  Email sent to {to}.")


def post_issue(title: str, body: str) -> None:
    token = os.environ.get("GITHUB_TOKEN", "").strip()
    repo = os.environ.get("GITHUB_REPOSITORY", "").strip()
    if not token or not repo:
        print("  Issue skipped (not running in GitHub Actions).")
        return
    api = f"https://api.github.com/repos/{repo}"
    h = {"Authorization": f"Bearer {token}", "Accept": "application/vnd.github+json"}
    old = requests.get(f"{api}/issues", headers=h,
                       params={"labels": "daily-echo", "state": "open"}, timeout=30).json()
    for issue in old if isinstance(old, list) else []:
        requests.patch(f"{api}/issues/{issue['number']}", headers=h,
                       json={"state": "closed"}, timeout=30)
    if len(body) > 65000:
        body = body[:65000] + "\n\n_(truncated \u2014 see the full report in `reports/`)_"
    r = requests.post(f"{api}/issues", headers=h,
                      json={"title": title, "body": body, "labels": ["daily-echo"]}, timeout=30)
    r.raise_for_status()
    print(f"  Issue created: {r.json().get('html_url')}")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-llm", action="store_true", help="only fetch and list headlines")
    args = ap.parse_args()

    day = datetime.now(TZ)
    print("Fetching headlines...")
    headlines = fetch_headlines()
    counts = {k: len(v) for k, v in headlines.items()}
    if args.no_llm:
        for cat, items in headlines.items():
            print(f"\n[{cat}]")
            for h in items[:5]:
                print(f"  - {h['title']}")
        return 0

    guidance = (ROOT / "guidance.md").read_text(encoding="utf-8")
    llm = LLM()
    stories: list[dict] = []
    for category, items in headlines.items():
        print(f"Analyzing {category}...")
        try:
            stories += analyze(llm, category, items, guidance)
        except Exception as exc:
            print(f"  ! analysis failed for {category}: {exc}")
    if not stories:
        print("No stories analyzed; aborting.")
        return 1

    print("Verifying precedents on Wikipedia...")
    verify(stories)

    md = render_markdown(day, stories, counts)
    REPORTS.mkdir(exist_ok=True)
    (REPORTS / f"{day:%Y-%m-%d}.md").write_text(md, encoding="utf-8")
    (REPORTS / "latest.md").write_text(md, encoding="utf-8")
    (DATA / "json").mkdir(parents=True, exist_ok=True)
    (DATA / "json" / f"{day:%Y-%m-%d}.json").write_text(
        json.dumps(stories, indent=2, ensure_ascii=False), encoding="utf-8")
    update_board(BOARD, day, stories)
    print(f"Saved report and board ({len(stories)} stories).")

    title = f"History Echo \u2014 {day:%a %b %d, %Y}"
    try:
        post_issue(title, md)
    except Exception as exc:
        print(f"  ! issue failed: {exc}")
    try:
        send_email(title, render_html(md), md)
    except Exception as exc:
        print(f"  ! email failed: {exc}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
